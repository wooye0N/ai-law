# -*- coding: utf-8 -*-
"""
桂法通 · 公共法律服务 AI 助手
运行： streamlit run ai_law.py
"""
from __future__ import annotations

import os
import re
import io
import time
from pathlib import Path
from typing import Optional, Tuple

import streamlit as st

# ============================================================
# 0. 页面配置 & 会话状态
# ============================================================
st.set_page_config(
    page_title="桂法通 · 公共法律服务 AI 助手",
    page_icon="⚖️",
    layout="wide",
    initial_sidebar_state="expanded",
)

st.markdown(
    """
    <style>
    html, body, [class*="css"] { font-size: 17px !important; }
    .stMarkdown p, .stMarkdown li { font-size: 17px !important; line-height: 1.75 !important; }
    section[data-testid="stSidebar"] .stMarkdown p,
    section[data-testid="stSidebar"] .stCaption { font-size: 15px !important; line-height: 1.6 !important; }
    button[data-baseweb="tab"] p { font-size: 17px !important; font-weight: 600 !important; }
    textarea, input, .stButton button, .stDownloadButton button { font-size: 16px !important; }
    h3, h4 { font-size: 21px !important; }
    </style>
    """,
    unsafe_allow_html=True,
)

DEFAULT_BASE_URL = "https://open.bigmodel.cn/api/paas/v4"
DEFAULT_MODEL = "glm-4-flash"
KB_DIR = Path(__file__).parent / "kb"
NPC_CACHE_DIR = Path(__file__).parent / "kb_npc_cache"
BUNDLE_GZ = Path(__file__).parent / "kb_bundle.jsonl.gz"
BUNDLE_JSONL = Path(__file__).parent / "kb_bundle.jsonl"

# 首次启动自动拉取的核心法律（覆盖民事 / 劳动 / 刑事 / 行政 / 消费维权）
CORE_LAWS = [
    "中华人民共和国民法典",
    "中华人民共和国劳动合同法",
    "中华人民共和国劳动争议调解仲裁法",
    "中华人民共和国法律援助法",
    "中华人民共和国人民调解法",
    "中华人民共和国仲裁法",
    "中华人民共和国刑法",
    "中华人民共和国治安管理处罚法",
    "中华人民共和国消费者权益保护法",
]

for _k, _v in {
    "api_key": os.getenv("LLM_API_KEY", ""),
    "base_url": os.getenv("LLM_BASE_URL", DEFAULT_BASE_URL),
    "model": os.getenv("LLM_MODEL", DEFAULT_MODEL),
}.items():
    st.session_state.setdefault(_k, _v)


# ============================================================
# 1. LLM 封装
# ============================================================
def _client():
    from openai import OpenAI
    return OpenAI(api_key=st.session_state["api_key"], base_url=st.session_state["base_url"], timeout=120.0)


def llm_stream(messages, temperature=0.3, max_tokens=2048):
    resp = _client().chat.completions.create(
        model=st.session_state["model"], messages=messages,
        temperature=temperature, max_tokens=max_tokens, stream=True,
    )
    for chunk in resp:
        try:
            delta = chunk.choices[0].delta.content
        except (IndexError, AttributeError):
            delta = None
        if delta:
            yield delta


def llm(messages, **kw):
    return "".join(llm_stream(messages, **kw))


# ============================================================
# 2. 国家法律法规数据库拉取
# ============================================================
NPC_BASE = "https://flk.npc.gov.cn"
NPC_HEADERS = {
    "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36",
    "Accept": "application/json, text/plain, */*",
    "Content-Type": "application/json",
}


def _npc_search_law(law_name: str) -> Optional[str]:
    """搜索现行有效法律，返回 bbbs id。"""
    import requests

    payload = {
        "searchRange": 1,
        "searchType": 1,
        "searchContent": law_name,
        "pageNum": 1,
        "pageSize": 5,
        "orderByParam": {"order": "-1", "sort": ""},
        "sxx": [3],
    }
    try:
        r = requests.post(
            f"{NPC_BASE}/law-search/search/list",
            headers=NPC_HEADERS, json=payload, timeout=15,
        )
        r.raise_for_status()
        data = r.json()
        rows = data.get("data", {}).get("list", []) or data.get("rows", [])
        if not rows:
            return None
        for row in rows:
            title = row.get("title", "") or row.get("bt", "")
            bbbs = row.get("bbbs") or row.get("id", "")
            if bbbs and (law_name in title or title in law_name):
                return bbbs
        first = rows[0]
        return first.get("bbbs") or first.get("id")
    except Exception as e:
        print(f"[NPC] 搜索失败 {law_name}: {e}")
        return None


def _npc_download_docx(bbbs: str, dest: Path) -> bool:
    """通过下载 API 获取 signed URL 并下载 docx。"""
    import requests

    try:
        r = requests.get(
            f"{NPC_BASE}/law-search/download/pc",
            params={"format": "docx", "bbbs": bbbs},
            headers=NPC_HEADERS, timeout=20,
        )
        r.raise_for_status()
        data = r.json()
        url = data.get("data", {}).get("url")
        if not url:
            print(f"[NPC] 无下载链接 bbbs={bbbs}")
            return False
        fr = requests.get(url, timeout=30)
        fr.raise_for_status()
        dest.parent.mkdir(parents=True, exist_ok=True)
        dest.write_bytes(fr.content)
        return True
    except Exception as e:
        print(f"[NPC] 下载失败 bbbs={bbbs}: {e}")
        return False


def fetch_law_from_npc(law_name: str) -> Optional[Path]:
    """拉取单部法律，返回本地 docx 路径；失败返回 None。"""
    safe_name = re.sub(r'[\\/:*?"<>|]', "_", law_name)
    cache_file = NPC_CACHE_DIR / f"{safe_name}.docx"
    if cache_file.exists() and cache_file.stat().st_size > 1000:
        return cache_file

    print(f"[NPC] 正在拉取：{law_name}")
    bbbs = _npc_search_law(law_name)
    if not bbbs:
        return None
    time.sleep(0.5)
    if _npc_download_docx(bbbs, cache_file):
        print(f"[NPC] 已保存：{cache_file.name}")
        return cache_file
    return None


# ============================================================
# 3. 法条知识库（本地 kb/ + 自动拉取 kb_npc_cache/）
# ============================================================
_ART_PAT = re.compile(r"^第[一二三四五六七八九十百千零〇0-9]+条")


def _split_articles(text: str):
    chunks, cur = [], ""
    for raw in text.splitlines():
        line = raw.strip()
        if not line:
            continue
        if _ART_PAT.match(line):
            if cur:
                chunks.append(cur)
            cur = line
        elif cur:
            cur = cur + " " + line
        else:
            cur = line
    if cur:
        chunks.append(cur)
    return chunks


def _read_docx_text(path: Path) -> str:
    from docx import Document
    doc = Document(str(path))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    parts.append(cell.text)
    return "\n".join(parts)


def _load_docx_articles(path: Path, source_name: str):
    try:
        raw = _read_docx_text(path)
    except Exception as e:
        print(f"[知识库] 读取 {path.name} 失败：{e}")
        return []
    return [{"text": art, "source": source_name} for art in _split_articles(raw) if len(art) >= 8]


def _load_bundle():
    """从打包好的 kb_bundle 读取法条（云端部署用，避免上传上千个小文件）。"""
    import gzip
    import json

    for p in (BUNDLE_GZ, BUNDLE_JSONL):
        if not p.exists():
            continue
        try:
            fh = gzip.open(p, "rt", encoding="utf-8") if p.suffix == ".gz" else p.open(encoding="utf-8")
            out = []
            with fh as f:
                for line in f:
                    line = line.strip()
                    if not line:
                        continue
                    obj = json.loads(line)
                    out.append({"text": obj["t"], "source": obj["s"]})
            print(f"[知识库] 已加载打包法条 {len(out)} 条（{p.name}）")
            return out
        except Exception as e:
            print(f"[知识库] 读取 {p.name} 失败：{e}")
    return []


@st.cache_resource(show_spinner="正在加载法条知识库（首次会自动从国家法律法规数据库拉取核心法律）…")
def load_kb():
    import jieba
    from rank_bm25 import BM25Okapi

    jieba.setLogLevel(60)
    docs = []

    # 0) 打包法条（优先，含广西地方性法规等扩展库）
    docs.extend(_load_bundle())

    # 1) 本地 kb/ 目录（txt + docx）
    if KB_DIR.exists():
        for f in sorted(list(KB_DIR.glob("*.txt")) + list(KB_DIR.glob("*.docx"))):
            try:
                if f.suffix.lower() == ".docx":
                    docs.extend(_load_docx_articles(f, f.stem))
                else:
                    raw = f.read_text(encoding="utf-8")
                    docs.extend([{"text": a, "source": f.stem} for a in _split_articles(raw) if len(a) >= 8])
            except Exception as e:
                print(f"[知识库] 跳过 {f.name}：{e}")

    # 2) 自动从 NPC 拉取核心法律（缓存优先）
    NPC_CACHE_DIR.mkdir(exist_ok=True)
    for law in CORE_LAWS:
        path = fetch_law_from_npc(law)
        if path:
            docs.extend(_load_docx_articles(path, path.stem))

    if not docs:
        return None

    # 去重（打包库与 kb/ 目录可能有重叠）
    seen, uniq = set(), []
    for d in docs:
        key = (d["source"], d["text"][:60])
        if key in seen:
            continue
        seen.add(key)
        uniq.append(d)
    docs = uniq

    corpus = [list(jieba.cut(d["text"])) for d in docs]
    return {"docs": docs, "bm25": BM25Okapi(corpus)}


# ---- 检索预处理：泛词剔除 + 主题词加权 ----
_STOPWORDS = {
    "注意", "需要", "什么", "怎么", "如何", "可以", "应该", "是否",
    "一个", "一些", "哪些", "这个", "那个", "我们", "你们", "他们",
    "问题", "情况", "时候", "现在", "一下", "咨询", "请问", "想要",
    "的话", "还是", "或者", "以及", "因为", "所以", "但是", "如果",
}

_TOPIC_WORDS = {
    # 民事 / 劳动
    "结婚", "婚姻", "离婚", "彩礼", "继承", "遗嘱",
    "欠钱", "借款", "债务", "利息", "还款",
    "工伤", "辞退", "试用期", "加班", "工资", "劳动合同",
    "租房", "租金", "押金", "买卖", "合同", "违约",
    "交通事故", "赔偿", "侵权", "抚养", "赡养",
    # 刑事
    "犯罪", "刑事", "判刑", "坐牢", "盗窃", "诈骗", "抢劫", "故意伤害",
    "醉驾", "酒驾", "赌博", "毒品", "走私", "贪污", "受贿", "正当防卫",
    # 行政
    "治安", "拘留", "罚款", "行政处罚", "扰乱", "斗殴", "寻衅滋事",
    # 消费维权
    "消费", "退货", "退款", "假货", "欺诈", "三包", "维权", "商家",
}


def _query_terms(query: str):
    import jieba
    words = [w for w in jieba.cut(query) if len(w) >= 2 and w not in _STOPWORDS]
    if not words:
        words = [w for w in jieba.cut(query) if len(w) >= 2]
    weighted = []
    for w in words:
        weighted.append(w)
        if w in _TOPIC_WORDS:
            weighted.extend([w, w])
    return weighted, set(words)


def retrieve(query: str, k: int = 8, min_score: float = 0.3, min_ratio: float = 0.2, min_overlap: int = 1):
    kb = load_kb()
    if not kb:
        return []
    terms, q_words = _query_terms(query)
    if not terms:
        return []
    scores = kb["bm25"].get_scores(terms)
    order = scores.argsort()[::-1]
    top = float(scores[order[0]]) if len(order) else 0.0
    if top <= 0:
        return []
    threshold = max(min_score, top * min_ratio)
    hits = []
    for i in order[:k]:
        s = float(scores[i])
        if s < threshold:
            break
        text = kb["docs"][i]["text"]
        overlap = sum(1 for w in q_words if w in text)
        if overlap < min_overlap:
            continue
        hits.append({**kb["docs"][i], "score": s})
    return hits


def kb_stats():
    kb = load_kb()
    if not kb:
        return 0, 0
    return len({d["source"] for d in kb["docs"]}), len(kb["docs"])


# ---- 从 AI 回答中提取法条引用并回查原文 ----
_CITE_PAT = re.compile(r"《([^》]+)》\s*第([一二三四五六七八九十百千零〇0-9]+)条")


def extract_citations(text: str):
    seen, out = set(), []
    for law, no in _CITE_PAT.findall(text):
        key = (law.strip(), no.strip())
        if key not in seen:
            seen.add(key)
            out.append(key)
    return out


def _norm_law_name(name: str) -> str:
    """去掉'中华人民共和国'前缀，便于回查匹配。"""
    return name.replace("中华人民共和国", "").strip()


def lookup_article(law_name: str, art_no: str) -> Optional[str]:
    kb = load_kb()
    if not kb:
        return None
    target = f"第{art_no}条"
    short = _norm_law_name(law_name)
    for d in kb["docs"]:
        src = d["source"]
        if (short in src or law_name in src) and d["text"].startswith(target):
            return d["text"]
    for d in kb["docs"]:
        src = d["source"]
        if (short in src or law_name in src) and target in d["text"][:40]:
            return d["text"]
    return None


# ============================================================
# 3.5 合同文件解析
# ============================================================
def read_contract_file(uploaded) -> Tuple[str, Optional[str]]:
    if uploaded is None:
        return "", None
    name = uploaded.name.lower()
    data = uploaded.read()

    if name.endswith(".txt"):
        for enc in ("utf-8", "utf-8-sig", "gbk", "gb18030"):
            try:
                return data.decode(enc), None
            except UnicodeDecodeError:
                continue
        return data.decode("utf-8", errors="ignore"), None

    if name.endswith(".docx"):
        try:
            from docx import Document
            doc = Document(io.BytesIO(data))
            parts = [p.text for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        if cell.text.strip():
                            parts.append(cell.text)
            text = "\n".join(parts)
            if not text.strip():
                return "", "这个 .docx 里没有读到文字，可能是纯图片扫描件。"
            return text, None
        except ImportError:
            return "", "未安装 python-docx，无法读取 .docx。"
        except Exception as e:
            return "", f"读取 .docx 失败：{e}"

    if name.endswith(".doc"):
        return "", "暂不支持老式 .doc 二进制格式。请用 Word 打开后「另存为」.docx，再上传。"

    return "", "仅支持 .txt 或 .docx 文件。"


# ============================================================
# 4. 模块一：智能法律咨询
# ============================================================
SYSTEM_QA = """你是「桂法通」，一名面向普通群众的公共法律服务 AI 助手，服务范围包括：法律咨询、法律援助、人民调解、仲裁、公证、司法鉴定等。

【工作规则】
1. 优先依据下面给出的【参考法条】作答；参考法条不足时，可依据通用法律常识补充，但引用法条必须准确。
2. 严禁编造法条条号或法条内容。引用法条时必须写明"《法律名称》第X条"。
3. 语言通俗、克制、共情，不吓唬当事人，也不做胜诉承诺。
4. 注意区分民事、刑事、行政责任。涉及刑事犯罪的，必须提示当事人及时咨询刑事律师，不要自行处理。

【输出格式】（Markdown）
**结论** —— 一到两句话直接回答核心问题。
**法律依据** —— 逐条列出引用的法条，格式必须是"《法律名称》第X条：要点"。
**你可以这样做** —— 3~5 条可执行的具体步骤。
**需要准备的材料** —— 如适用。
**注意时效与风险** —— 如适用（诉讼时效、举证责任、刑事风险等）。
**温馨提示** —— 固定一句：以上为一般性法律信息参考，不构成正式法律意见；具体案件建议咨询执业律师或拨打 12348。
"""


def render_qa():
    st.markdown("#### 💬 智能法律咨询")
    st.caption("描述你遇到的法律问题，系统会检索现行法条并给出可执行的建议。")

    q = st.text_area("你的问题", height=110,
                     placeholder="例：公司试用期说好 3 个月，结果一直没转正，也没签书面合同，我该怎么办？",
                     key="qa_input")
    go = st.button("开始解答", type="primary", key="qa_go")

    if not go or not q.strip():
        return
    if not st.session_state["api_key"]:
        st.error("请先在左侧填写 API Key。")
        return

    with st.spinner("正在检索相关法条…"):
        hits = retrieve(q, k=8)

    if hits:
        with st.expander(f"📚 检索到 {len(hits)} 条候选法条（供 AI 参考）"):
            for h in hits:
                snippet = h["text"][:150] + "…" if len(h["text"]) > 150 else h["text"]
                st.markdown(f"- **{h['source']}**：{snippet}")
        ctx = "\n".join(f"【{h['source']}】{h['text']}" for h in hits)
    else:
        st.caption("💡 未检索到候选法条，以下回答主要基于通用法律常识，供参考。")
        ctx = "（未匹配到法条）"

    messages = [
        {"role": "system", "content": SYSTEM_QA},
        {"role": "user", "content": f"【参考法条】\n{ctx}\n\n【当事人问题】\n{q}"},
    ]

    st.markdown("---")
    placeholder = st.empty()
    full_text = ""
    for chunk in llm_stream(messages, temperature=0.3, max_tokens=2400):
        full_text += chunk
        placeholder.markdown(full_text)

    cites = extract_citations(full_text)
    if cites:
        with st.expander(f"📚 AI 引用的 {len(cites)} 条法条原文（点击查看）"):
            for law, no in cites:
                original = lookup_article(law, no)
                if original:
                    snippet = original[:220] + "…" if len(original) > 220 else original
                    st.markdown(f"- **《{law}》第{no}条**：{snippet}")
                else:
                    st.markdown(f"- **《{law}》第{no}条**：⚠️ 知识库未收录该条原文，请以官方文本为准")
    else:
        st.caption("💡 本次回答未引用具体法条条号。")


# ============================================================
# 5. 模块二：法律文书生成
# ============================================================
DOC_TYPES = {
    "民事起诉状": ["原告信息（姓名 / 性别 / 出生日期 / 住址 / 联系电话）", "被告信息（姓名 / 住址 / 联系电话）", "诉讼请求（希望法院判什么）", "事实与理由（事情经过）", "受理法院"],
    "劳动仲裁申请书": ["申请人信息（姓名 / 身份证号 / 住址 / 电话）", "被申请人信息（公司全称 / 地址 / 法定代表人）", "仲裁请求", "事实与理由", "受理仲裁委员会"],
    "法律援助申请书": ["申请人信息（姓名 / 身份证号 / 住址 / 电话）", "家庭经济状况", "申请事项（需要什么帮助）", "事实与理由", "受理机构（XX 法律援助中心）"],
    "人民调解协议书": ["甲方信息", "乙方信息", "纠纷简要情况", "达成协议的内容", "履行方式与期限", "调解委员会名称"],
    "借条": ["出借人姓名", "借款人姓名 / 身份证号", "借款金额（大写 + 小写）", "借款日期", "还款日期", "利息约定", "交付方式（现金 / 转账）"],
}

SYSTEM_DOC = """你是法律文书起草助手。请根据用户提供的信息，起草一份规范、完整、可直接使用的【{doctype}】。

【要求】
1. 使用标准法律文书格式：标题、当事人信息、正文、结尾（致送机构 / 署名 / 日期）。
2. 只使用用户提供的信息。缺失的关键信息用【请补充：xxx】占位，绝对不要编造姓名、金额、日期。
3. 语言严谨、客观、书面化，不使用情绪化表述。
4. 直接输出文书正文，不要额外解释或寒暄。
5. 文书末尾另起一行，加一句小字提示：本文书由 AI 辅助生成，正式提交前请核对事实并咨询专业人士。
"""


def to_docx(text: str, title: str):
    try:
        from docx import Document
        from docx.shared import Pt
    except ImportError:
        return None
    doc = Document()
    for line in text.splitlines():
        line = line.rstrip()
        if not line:
            continue
        p = doc.add_paragraph(line)
        p.paragraph_format.space_after = Pt(0)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def render_doc():
    st.markdown("#### 📝 法律文书生成")
    st.caption("填写关键信息，一键生成格式规范、可直接使用的法律文书。")
    doctype = st.selectbox("文书类型", list(DOC_TYPES.keys()))
    fields = DOC_TYPES[doctype]

    with st.form("doc_form"):
        values = {}
        for f in fields:
            values[f] = st.text_area(f, height=70, key=f"doc_{doctype}_{f}")
        submitted = st.form_submit_button("生成文书", type="primary")

    if not submitted:
        return
    if not st.session_state["api_key"]:
        st.error("请先在左侧填写 API Key。")
        return

    info = "\n".join(f"- {k}：{v.strip() or '（未提供）'}" for k, v in values.items())
    messages = [
        {"role": "system", "content": SYSTEM_DOC.replace("{doctype}", doctype)},
        {"role": "user", "content": f"请起草一份【{doctype}】。\n当事人提供的信息：\n{info}"},
    ]

    st.markdown("---")
    with st.spinner("正在起草…"):
        text = llm(messages, temperature=0.2, max_tokens=2400)
    st.markdown(text)

    c1, c2 = st.columns(2)
    c1.download_button("⬇️ 下载 .txt", text.encode("utf-8"), file_name=f"{doctype}.txt", mime="text/plain")
    docx_bytes = to_docx(text, doctype)
    if docx_bytes:
        c2.download_button("⬇️ 下载 .docx", docx_bytes, file_name=f"{doctype}.docx",
                           mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document")


# ============================================================
# 6. 模块三：合同智能体检
# ============================================================
CONTRACT_CHECKLIST = {
    "合同主体": ["甲方", "乙方"], "标的与价款": ["金额", "价款", "报酬", "费用"],
    "履行期限": ["期限", "之日起", "日内", "之前"], "违约责任": ["违约"],
    "争议解决": ["争议", "诉讼", "仲裁"], "合同解除": ["解除"],
    "不可抗力": ["不可抗力"], "签署要件": ["签字", "盖章", "签署"],
}

CONTRACT_RISKS = [
    (r"最终解释权", "「最终解释权归本方所有」属于排除对方权利的格式条款，可能被认定无效。"),
    (r"概不(负责|承担)|不承担任何(责任|损失)", "绝对免责表述。依《民法典》第 506 条，造成人身损害或因故意/重大过失致财产损失的免责条款无效。"),
    (r"自动(续约|延期|展期)", "自动续约条款，注意是否约定提前通知解除的窗口期。"),
    (r"违约金.{0,20}?\d{2,3}\s*%", "违约金比例偏高，超过实际损失 30% 的部分可能被法院调减（《民法典》第 585 条）。"),
    (r"定金.{0,15}?\d{2,3}\s*%", "定金不得超过主合同标的额的 20%，超出部分不产生定金效力。"),
    (r"管辖.{0,20}(外地|异地|对方所在地)", "管辖约定涉及异地诉讼成本，需评估是否属于格式条款加重负担。"),
    (r"视为(认可|同意|接受)|无异议", "默示同意条款，注意是否留了合理异议期。"),
    (r"押金|保证金", "押金 / 保证金的退还条件与期限是否明确？"),
]

SYSTEM_CONTRACT = """你是一名合同审查律师助理。请基于用户提供的合同文本，输出一份简洁的《合同体检报告》。

【输出格式】（Markdown）
## 一、整体评价
一句话说明这份合同的主要风险等级（高 / 中 / 低）及原因。

## 二、重点风险条款
逐条列出，每条包含：**原文摘录** → **风险说明** → **修改建议（给出可直接替换的措辞）**。最多 6 条，按严重程度排序。

## 三、缺失条款提示
指出合同缺少的常见保护性条款，并说明为什么需要补充。

## 四、签署前提醒
3~4 条可执行动作。

【硬性要求】
- 只针对用户给的文本分析，不要虚构合同里没有的内容。
- 语言简明，避免长篇大论。
- 结尾固定一句：本报告为 AI 辅助审查，重大合同请交由执业律师复核。
"""


def rule_check(text: str):
    missing = [name for name, kws in CONTRACT_CHECKLIST.items() if not any(kw in text for kw in kws)]
    risks = []
    for pat, desc in CONTRACT_RISKS:
        m = re.search(pat, text)
        if m:
            s = max(0, m.start() - 25)
            e = min(len(text), m.end() + 25)
            risks.append({"desc": desc, "snippet": text[s:e].replace("\n", " ").strip()})
    return missing, risks


def render_contract():
    st.markdown("#### 🔍 合同智能体检")
    st.caption("支持上传 .txt / .docx，或直接粘贴合同文本。先跑规则引擎扫明显坑，再由 AI 出整体报告。")

    text = ""
    up = st.file_uploader("上传合同（.txt 或 .docx）", type=["txt", "docx", "doc"],
                          help=".doc 为老式二进制格式，请先用 Word「另存为」.docx 再上传。")
    if up is not None:
        file_text, err = read_contract_file(up)
        if err:
            st.error(err)
        elif file_text.strip():
            text = file_text
            st.success(f"已读取「{up.name}」，共 {len(text)} 字。")
        else:
            st.warning("文件内容为空。")

    pasted = st.text_area("或直接粘贴合同文本", height=220, placeholder="把合同全文粘在这里…", key="contract_text")
    if pasted.strip():
        text = pasted

    if st.button("开始体检", type="primary", key="contract_go"):
        if not text.strip():
            st.warning("请先上传或粘贴合同文本。")
            return

        missing, risks = rule_check(text)
        st.markdown("### 📋 规则引擎快扫")
        c1, c2 = st.columns(2)
        with c1:
            if missing:
                st.error("**可能缺失的条款**\n\n" + "\n".join(f"- {m}" for m in missing))
            else:
                st.success("**条款结构完整**，常见必备条款均已出现。")
        with c2:
            if risks:
                st.warning("**命中的风险表述**\n\n" + "\n".join(f"- {r['desc']}\n  > 原文：…{r['snippet']}…" for r in risks))
            else:
                st.success("**未命中已知风险模式。**")

        if not st.session_state["api_key"]:
            st.info("提示：填写左侧 API Key 可获得 AI 深度分析报告。")
            return

        st.markdown("---")
        st.markdown("### 🤖 AI 深度体检报告")
        messages = [
            {"role": "system", "content": SYSTEM_CONTRACT},
            {"role": "user", "content": f"【合同文本】\n{text[:12000]}"},
        ]
        st.write_stream(llm_stream(messages, temperature=0.2, max_tokens=2600))


# ============================================================
# 7. 主入口
# ============================================================
def main():
    with st.sidebar:
        st.markdown("### ⚙️ 模型设置")
        st.text_input("API Key", type="password", key="api_key")
        st.text_input("Base URL", key="base_url")
        st.text_input("模型名称", key="model")
        st.caption("推荐免费方案：\n- 智谱 GLM-4-Flash（open.bigmodel.cn）\n- 硅基流动 Qwen2.5-7B（siliconflow.cn）\n- 本地 Ollama：base_url=http://localhost:11434/v1")

        n_law, n_art = kb_stats()
        st.markdown("---")
        st.markdown(f"### 📚 法条知识库\n已加载 **{n_law}** 部法规 / **{n_art}** 条法条")
        st.caption("覆盖：民事 / 劳动 / 刑事 / 行政 / 消费维权。本地 kb/ + 自动拉取缓存 kb_npc_cache/。")

        if st.button("🔄 刷新知识库"):
            st.cache_resource.clear()
            st.rerun()

        if n_art == 0:
            st.warning("知识库为空。请检查网络或手动在 kb/ 放入法规文件。")

        st.markdown("---")
        st.caption("⚠️ 本工具输出为 AI 辅助信息，不构成正式法律意见。")

    st.title("⚖️ 桂法通")
    st.markdown("##### 面向公共法律服务的 AI 助手 · 法律咨询 / 文书生成 / 合同体检")

    t1, t2, t3 = st.tabs(["💬 法律咨询", "📝 文书生成", "🔍 合同体检"])
    with t1:
        render_qa()
    with t2:
        render_doc()
    with t3:
        render_contract()


if __name__ == "__main__":
    main()